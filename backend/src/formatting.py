#!/usr/bin/env python3
"""
Утилиты для конвертации Markdown (.md) в DOCX.
Используется пайплайном для генерации финального отчёта.
"""

import re
from urllib.parse import urlparse

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def extract_domain(url: str) -> str:
    """Извлекает домен из URL с обрезкой длины."""
    try:
        parsed = urlparse(url)
        domain = parsed.netloc
        if domain.startswith("www."):
            domain = domain[4:]
        if len(domain) > 30:
            domain = domain[:27] + "..."
        return domain
    except Exception:
        return url[:30] + "..." if len(url) > 30 else url


def add_hyperlink(paragraph, text: str, url: str) -> None:
    """Добавляет гиперссылку в параграф."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0000FF")
    rPr.append(color)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.append(underline)

    run.append(rPr)

    text_elem = OxmlElement("w:t")
    text_elem.text = text
    run.append(text_elem)

    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_horizontal_line(doc: Document) -> None:
    """Добавляет горизонтальную линию."""
    paragraph = doc.add_paragraph()
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    pBdr.append(bottom)
    pPr.append(pBdr)


def process_text_with_formatting(text: str) -> list[tuple[str, str, str | None]]:
    """
    Обрабатывает строку Markdown-текста.
    Возвращает список (тип, содержимое, url).
    """
    if not text or text == "---":
        return [("text", text, None)]

    result: list[tuple[str, str, str | None]] = []
    i = 0
    length = len(text)

    while i < length:
        url_match = re.match(r'https?://[^\s<>"{}|\\^`\[\]]+', text[i:])
        if url_match:
            url = url_match.group(0)
            result.append(("url", url, url))
            i += len(url)
            continue

        if i + 1 < length and text[i : i + 2] == "**":
            end = text.find("**", i + 2)
            if end != -1:
                content = text[i + 2 : end]
                result.append(("bold", content, None))
                i = end + 2
                continue
            else:
                result.append(("text", text[i], None))
                i += 1
                continue

        if text[i] == "*" and (i + 1 >= length or text[i + 1] != "*"):
            end = text.find("*", i + 1)
            if end != -1:
                content = text[i + 1 : end]
                result.append(("italic", content, None))
                i = end + 1
                continue
            else:
                result.append(("text", text[i], None))
                i += 1
                continue

        if text[i] == "[":
            bracket_end = text.find("]", i)
            if bracket_end != -1 and bracket_end + 1 < length and text[bracket_end + 1] == "(":
                paren_end = text.find(")", bracket_end + 2)
                if paren_end != -1:
                    link_text = text[i + 1 : bracket_end]
                    url = text[bracket_end + 2 : paren_end]
                    result.append(("link", link_text, url))
                    i = paren_end + 1
                    continue
                else:
                    result.append(("text", text[i], None))
                    i += 1
                    continue
            else:
                result.append(("text", text[i], None))
                i += 1
                continue

        start = i
        while i < length:
            current_char = text[i]
            if (
                current_char == "*"
                or current_char == "["
                or (i + 1 < length and text[i : i + 2] == "**")
                or (i + 7 < length and text[i : i + 7] == "http://")
                or (i + 8 < length and text[i : i + 8] == "https://")
            ):
                break
            i += 1

        if i > start:
            result.append(("text", text[start:i], None))

    return result


def add_formatted_text_to_paragraph(paragraph, text: str | None) -> None:
    """Добавляет форматированный текст в параграф."""
    if text is None:
        return

    if text == "---" or (text.startswith(":---") and text.endswith("---:")):
        return

    parts = process_text_with_formatting(text)

    for part_type, content, url in parts:
        if part_type == "text":
            if content and content != "---":
                paragraph.add_run(content)
        elif part_type == "bold":
            run = paragraph.add_run(content)
            run.bold = True
        elif part_type == "italic":
            run = paragraph.add_run(content)
            run.italic = True
        elif part_type == "link":
            add_hyperlink(paragraph, content, url or "")
        elif part_type == "url":
            domain = extract_domain(content)
            add_hyperlink(paragraph, domain, content)


def convert_md_to_docx(input_file: str, output_file: str) -> None:
    """Конвертирует Markdown в DOCX."""
    with open(input_file, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()
    i = 0

    while i < len(lines):
        line = lines[i].rstrip()

        if "|" in line and i + 1 < len(lines) and "|" in lines[i + 1]:
            table_lines: list[str] = []
            while i < len(lines) and "|" in lines[i]:
                table_lines.append(lines[i].rstrip())
                i += 1

            if len(table_lines) >= 2:
                rows: list[list[str]] = []
                for table_line in table_lines:
                    cells = [cell.strip() for cell in table_line.split("|")]
                    if cells and not cells[0]:
                        cells = cells[1:]
                    if cells and not cells[-1]:
                        cells = cells[:-1]
                    rows.append(cells)

                header_row = rows[0] if rows else []
                data_rows: list[list[str]] = []

                for row in rows[1:]:
                    if row and all(re.match(r"^:?-{3,}:?$", cell) for cell in row):
                        continue
                    data_rows.append(row)

                if header_row:
                    cols = len(header_row)
                    for row in data_rows:
                        cols = max(cols, len(row))

                    table = doc.add_table(rows=1 + len(data_rows), cols=cols)
                    table.style = "Table Grid"

                    for j, cell_text in enumerate(header_row):
                        if j < cols:
                            cell = table.cell(0, j)
                            cell.text = ""
                            paragraph = cell.paragraphs[0]
                            add_formatted_text_to_paragraph(paragraph, cell_text)
                            for run in paragraph.runs:
                                run.bold = True

                    for row_idx, row_data in enumerate(data_rows):
                        for col_idx, cell_text in enumerate(row_data):
                            if col_idx < cols:
                                cell = table.cell(row_idx + 1, col_idx)
                                cell.text = ""
                                paragraph = cell.paragraphs[0]
                                add_formatted_text_to_paragraph(paragraph, cell_text)

            continue

        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", line):
            add_horizontal_line(doc)
            i += 1
            continue

        if line.startswith("#"):
            level = len(re.match(r"^#+", line).group())
            text = line.lstrip("#").strip()

            heading = doc.add_heading(level=min(level, 4))
            if text:
                add_formatted_text_to_paragraph(heading, text)

            i += 1
            continue

        list_match = re.match(r"^([\-\*]|\d+\.)\s+(.*)", line)
        if list_match:
            list_type = list_match.group(1)
            content = list_match.group(2)

            if list_type in ["-", "*"]:
                paragraph = doc.add_paragraph(style="List Bullet")
            else:
                paragraph = doc.add_paragraph(style="List Number")

            if content:
                add_formatted_text_to_paragraph(paragraph, content)

            i += 1
            continue

        if line:
            paragraph = doc.add_paragraph()
            add_formatted_text_to_paragraph(paragraph, line)
        else:
            doc.add_paragraph()

        i += 1

    doc.save(output_file)

