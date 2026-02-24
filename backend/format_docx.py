import docx
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn

def set_font(run, font_name):
    """
    Надежно устанавливает шрифт для запуска (run), 
    включая настройки для латиницы и восточноазиатских символов.
    """
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run._element.rPr.rFonts.set(qn('w:ascii'), font_name)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), font_name)
    run._element.rPr.rFonts.set(qn('w:cs'), font_name)

def format_paragraph(paragraph, heading_styles):
    """
    Применяет форматирование к абзацу.
    Сначала весь текст на 14 кегль, затем заголовки получают свои размеры.
    """
    # 1. Выравнивание по ширине
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # 2. Межстрочный интервал 1.5
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    paragraph.paragraph_format.line_spacing = 1.5
    
    style_name = paragraph.style.name

    for run in paragraph.runs:
        # 3. Шрифт Times New Roman
        set_font(run, 'Times New Roman')

        # 4. Сначала устанавливаем всем текст 14 кегль (как Compact)
        run.font.size = Pt(14)

        # 5. Затем переопределяем размер для заголовков
        if style_name in heading_styles:
            run.font.size = Pt(heading_styles[style_name])

def format_docx(input_path, output_path):
    doc = Document(input_path)

    # Размеры шрифта только для заголовков (остальное всё 14 кегль по умолчанию)
    heading_styles = {
        'Heading 1': 20,
        'Заголовок 1': 20,
        'Heading 2': 16,
        'Заголовок 2': 16,
        'Heading 3': 14,
        'Заголовок 3': 14,
    }

    # 1. Форматируем основные абзацы документа
    for paragraph in doc.paragraphs:
        format_paragraph(paragraph, heading_styles)

    # 2. Форматируем абзацы внутри всех таблиц
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    format_paragraph(paragraph, heading_styles)

    doc.save(output_path)
    print(f"Документ успешно отформатирован и сохранен как {output_path}")

if __name__ == "__main__":
    input_file = 'FlatMapAI_DeepSeek_v3_1.docx'   
    output_file = 'formatted.docx' 
    
    try:
        format_docx(input_file, output_file)
    except FileNotFoundError:
        print(f"Ошибка: Файл '{input_file}' не найден.")
    except Exception as e:
        print(f"Произошла ошибка: {e}")