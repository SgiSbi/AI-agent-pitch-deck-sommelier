import sys
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn


# --- КОНФИГУРАЦИЯ ---
SOURCE_DIR = Path('tmp')
DEST_DIR = Path('result')

HEADING_STYLES = {
    'Heading 1': 20,
    'Заголовок 1': 20,
    'Heading 2': 16,
    'Заголовок 2': 16,
    'Heading 3': 14,
    'Заголовок 3': 14,
}
DEFAULT_FONT_SIZE = Pt(14)
FONT_NAME = 'Times New Roman'

def set_font(run, font_name):
    """
    Надежно устанавливает шрифт для запуска (run), 
    включая настройки для латиницы и восточноазиатских символов.
    """
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run._element.rPr.rFonts.set(qn('w:ascii'), font_name)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), font_name)
    run._element.rPr.rFonts.set(qn('w:cs'), font_name)

def format_paragraph(paragraph, heading_styles):
    """
    Применяет форматирование к абзацу.
    """
    # 1. Выравнивание по ширине
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # 2. Межстрочный интервал 1.5
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    paragraph.paragraph_format.line_spacing = 1.5
    
    style_name = paragraph.style.name

    for run in paragraph.runs:
        # 3. Шрифт Times New Roman
        set_font(run, FONT_NAME)

        # 4. Сначала устанавливаем всем текст 14 кегль
        run.font.size = DEFAULT_FONT_SIZE

        # 5. Затем переопределяем размер для заголовков
        if style_name in heading_styles:
            run.font.size = Pt(heading_styles[style_name])

def format_docx_file(input_path, output_path):
    """
    Открывает DOCX, применяет стили и сохраняет.
    """
    doc = Document(input_path)

    # 1. Форматируем основные абзацы документа
    for paragraph in doc.paragraphs:
        format_paragraph(paragraph, HEADING_STYLES)

    # 2. Форматируем абзацы внутри всех таблиц
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    format_paragraph(paragraph, HEADING_STYLES)

    doc.save(output_path)

def convert_md_to_docx(input_file: Path, output_file: Path) -> None:
    """
    Конвертация Markdown-файла в DOCX без.

    Поддерживаются базовые конструкции:
    - заголовки: #, ##, ###
    - маркированные списки: -, * или + в начале строки
    - нумерованные списки: "1. ", "2. " и т.п.
    - таблицы в markdown-формате (pipe-таблицы)
    - обычные абзацы
    Также переносит базовое inline-форматирование: **bold**, *italic*.
    Горизонтальные линии вида '---' не отображаются как текст.
    """
    doc = Document()

    inline_pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|__[^_]+__|_[^_]+_)")
    LIST_INDENT_STEP_PT = 18  # размер отступа для каждого уровня списка (в пунктах)

    def apply_inline_markdown(paragraph, text: str) -> None:
        """
        Заполняет абзац текстом с учётом простого inline-форматирования Markdown.
        """
        pos = 0
        for match in inline_pattern.finditer(text):
            start, end = match.span()
            if start > pos:
                paragraph.add_run(text[pos:start])

            token = match.group(0)
            if token.startswith("**") or token.startswith("__"):
                content = token[2:-2]
                run = paragraph.add_run(content)
                run.bold = True
            else:
                content = token[1:-1]
                run = paragraph.add_run(content)
                run.italic = True

            pos = end

        if pos < len(text):
            paragraph.add_run(text[pos:])

    def add_markdown_paragraph(text: str, style_name: str | None = None) -> None:
        p = doc.add_paragraph()
        if style_name:
            p.style = style_name
        apply_inline_markdown(p, text)

    def add_list_item(text: str, numbered: bool, indent_level: int) -> None:
        """
        Добавляет элемент списка с учётом уровня вложенности.
        indent_level считается в "ступенях" (0 = базовый уровень).
        """
        style_name = "List Number" if numbered else "List Bullet"
        p = doc.add_paragraph(style=style_name)
        if indent_level > 0:
            p.paragraph_format.left_indent = Pt(LIST_INDENT_STEP_PT * indent_level)
        apply_inline_markdown(p, text)

    def add_horizontal_rule() -> None:
        """
        Добавляет в документ горизонтальный разделитель (нижняя граница абзаца).
        """
        from docx.oxml import OxmlElement

        p = doc.add_paragraph()
        p_fmt = p.paragraph_format
        p_fmt.space_before = Pt(6)
        p_fmt.space_after = Pt(6)

        p_pr = p._p.get_or_add_pPr()
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is None:
            p_bdr = OxmlElement("w:pBdr")
            p_pr.append(p_bdr)

        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "auto")
        p_bdr.append(bottom)

    def is_table_separator(line: str) -> bool:
        """
        Проверяет, является ли строка разделителем таблицы вида:
        |---|---| или --- | :--- | ---:
        """
        if "|" not in line or "-" not in line:
            return False
        # убираем крайние пайпы
        core = line.strip()
        if core.startswith("|"):
            core = core[1:]
        if core.endswith("|"):
            core = core[:-1]
        # каждая ячейка должна состоять только из - : и пробелов
        for cell in core.split("|"):
            if not cell.strip():
                continue
            if not re.fullmatch(r":?-+:?", cell.strip()):
                return False
        return True

    def parse_table(start_index: int, all_lines: list[str]) -> tuple[list[list[str]], int]:
        """
        Разбирает markdown-таблицу, начиная с строки заголовка.
        Возвращает (rows, next_index):
        - rows: список строк таблицы, каждая строка — список ячеек
        - next_index: индекс следующей строки после таблицы
        """
        rows: list[list[str]] = []
        i = start_index
        n = len(all_lines)

        # первая строка — заголовок
        header = all_lines[i].strip()
        if not header or "|" not in header or i + 1 >= n:
            return rows, i

        sep = all_lines[i + 1].strip()
        if not is_table_separator(sep):
            return rows, i

        def split_row(row_line: str) -> list[str]:
            core = row_line.strip()
            if core.startswith("|"):
                core = core[1:]
            if core.endswith("|"):
                core = core[:-1]
            return [cell.strip() for cell in core.split("|")]

        header_cells = split_row(header)
        rows.append(header_cells)
        i += 2  # перескакиваем разделитель

        # последующие строки таблицы
        while i < n:
            line = all_lines[i].rstrip("\n")
            stripped = line.strip()
            if not stripped or "|" not in stripped:
                break
            rows.append(split_row(stripped))
            i += 1

        return rows, i

    def add_table_to_doc(rows: list[list[str]]) -> None:
        """
        Добавляет таблицу в документ, используя данные rows.
        Первая строка считается заголовком.
        """
        if not rows:
            return
        num_cols = max(len(r) for r in rows)
        table = doc.add_table(rows=len(rows), cols=num_cols)
        table.style = "Table Grid"

        for r_idx, row in enumerate(rows):
            for c_idx in range(num_cols):
                text = row[c_idx] if c_idx < len(row) else ""
                cell = table.rows[r_idx].cells[c_idx]
                # очищаем параграф по умолчанию
                cell_par = cell.paragraphs[0]
                cell_par.text = ""
                apply_inline_markdown(cell_par, text)
                if r_idx == 0:
                    for run in cell_par.runs:
                        run.bold = True

    with input_file.open("r", encoding="utf-8") as f:
        lines = f.readlines()

    last_was_blank = False
    i = 0
    n = len(lines)

    while i < n:
        raw_line = lines[i]
        line = raw_line.rstrip("\n")
        stripped = line.strip()

        # Пустая строка — просто учитываем как "перед нами была пустая строка"
        if not stripped:
            last_was_blank = True
            i += 1
            continue

        # Таблицы (pipe-таблицы)
        if "|" in stripped and i + 1 < n and is_table_separator(lines[i + 1].strip()):
            rows, next_i = parse_table(i, lines)
            if rows:
                add_table_to_doc(rows)
                # После таблицы добавляем один пустой абзац-разделитель,
                # чтобы отделить её от последующего текста/заголовка.
                doc.add_paragraph("")
                last_was_blank = True
                i = next_i
                continue

        # Горизонтальная линия (---, *** и т.п.) — рисуем реальную линию
        if re.fullmatch(r"[-*_]{3,}", stripped):
            add_horizontal_rule()
            last_was_blank = False
            i += 1
            continue

        # Заголовки
        if stripped.startswith("### "):
            # Пустая строка ПЕРЕД заголовком (если её ещё нет)
            if not last_was_blank:
                doc.add_paragraph("")
            add_markdown_paragraph(stripped[4:].strip(), style_name="Heading 3")
            # Пустая строка ПОСЛЕ заголовка
            doc.add_paragraph("")
            last_was_blank = True
            i += 1
            continue
        if stripped.startswith("## "):
            # Пустая строка ПЕРЕД заголовком (если её ещё нет)
            if not last_was_blank:
                doc.add_paragraph("")
            add_markdown_paragraph(stripped[3:].strip(), style_name="Heading 2")
            # Пустая строка ПОСЛЕ заголовка
            doc.add_paragraph("")
            last_was_blank = True
            i += 1
            continue
        if stripped.startswith("# "):
            # Пустая строка ПЕРЕД заголовком (если её ещё нет)
            if not last_was_blank:
                doc.add_paragraph("")
            add_markdown_paragraph(stripped[2:].strip(), style_name="Heading 1")
            # Пустая строка ПОСЛЕ заголовка
            doc.add_paragraph("")
            last_was_blank = True
            i += 1
            continue

        # Маркированный список (поддержка вложенности по количеству ведущих пробелов)
        m_bullet = re.match(r"^(\s*)([-*+])\s+(.+)$", line)
        if m_bullet:
            indent_spaces, _bullet, item_text = m_bullet.groups()
            level = max(len(indent_spaces) // 2, 0)
            add_list_item(item_text.strip(), numbered=False, indent_level=level)
            last_was_blank = False
            i += 1
            continue

        # Нумерованный список (поддержка вложенности по количеству ведущих пробелов)
        m_num = re.match(r"^(\s*)(\d+)\.\s+(.+)$", line)
        if m_num:
            indent_spaces, _num, item_text = m_num.groups()
            level = max(len(indent_spaces) // 2, 0)
            add_list_item(item_text.strip(), numbered=True, indent_level=level)
            last_was_blank = False
            i += 1
            continue

        # Обычный абзац
        add_markdown_paragraph(stripped)
        last_was_blank = False
        i += 1

    doc.save(str(output_file))

def process_files():
    """Основная логика сканирования и обработки."""

    # 1. Подготовка директорий
    if not SOURCE_DIR.exists():
        print(f"Ошибка: Исходная директория '{SOURCE_DIR}' не существует.")
        print("Создайте папку 'tmp' и положите туда .md файлы.")
        sys.exit(1)
    
    try:
        DEST_DIR.mkdir(parents=True, exist_ok=True)
    except PermissionError:
        print(f"Ошибка: Нет прав на создание директории '{DEST_DIR}'.")
        sys.exit(1)

    # 3. Поиск файлов
    md_files = list(SOURCE_DIR.glob('*.md'))
    
    if not md_files:
        print(f"В директории '{SOURCE_DIR}' не найдено файлов .md")
        sys.exit(0)

    print(f"Найдено файлов: {len(md_files)}")
    print(f"Начинаем обработку...\n")

    success_count = 0
    fail_count = 0

    for md_path in md_files:
        # Сохраняем с тем же именем, но с расширением .docx
        docx_path = DEST_DIR / f"{md_path.stem}.docx"
        
        print(f"Обработка: {md_path.name}...", end=" ")
        
        try:
            # Шаг 1: Конвертация
            convert_md_to_docx(md_path, docx_path)
            
            # Шаг 2: Форматирование (в тот же файл)
            format_docx_file(docx_path, docx_path)
            
            print("OK")
            success_count += 1
            
        except Exception as e:
            print(f"FAILED ({e})")
            fail_count += 1
            # Очищаем битый файл, если он создался
            if docx_path.exists():
                try:
                    docx_path.unlink()
                except:
                    pass

    print("\n" + "="*30)
    print(f"Готово. Успешно: {success_count}, Ошибок: {fail_count}")
    print(f"Результаты сохранены в: {DEST_DIR.resolve()}")

if __name__ == "__main__":
    process_files()